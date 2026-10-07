"""Daily per-symbol verdict workbook: direction (buy/sell/neutral) and confidence for every symbol traded today.
Confidence = estimated probability that price moves 5% in the signal's direction before 5% against it within 10 sessions.
Evidence: 5% model (calibrated out-of-sample on 1404-1405) + every rule that fired today for the symbol, each weighted by
how that rule has worked on that symbol (symbol_methods.csv: blended historical/live hit rate)."""
import os, sys, pickle, numpy as np, pandas as pd, live
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
BASE = os.path.dirname(os.path.abspath(__file__)); ST = os.path.join(BASE, 'live_state')
X, N = pd.read_pickle(os.path.join(BASE, 'panel.pkl'))
DAY = max(x.index[-1] for x in X.values())
day = DAY.date().isoformat()
# 1) model probabilities
fs = os.path.join(ST, 'five_scores.csv')
S = pd.read_csv(fs) if os.path.exists(fs) else None
if S is None or str(S.date.max())[:10] != day:
    import lightgbm as lgb
    g = {}; exec(open(os.path.join(BASE, 'five_feats.py')).read(), g)
    F, FEAT = g['F'], g['FEAT']; M = pickle.load(open(os.path.join(BASE, 'five_models.pkl'), 'rb'))
    S = F[F.date == F.date.max()].copy()
    for side in ['yb', 'ys']: S['p_' + side] = lgb.Booster(model_str=M['models'][side]).predict(S[FEAT])
    S = S[['sym', 'date', 'p_yb', 'p_ys']]; S.to_csv(fs, index=False)
S['sym'] = S.sym.map(live.norm); S = S.drop_duplicates('sym').set_index('sym')
CAL = pickle.load(open(os.path.join(BASE, 'five_calib.pkl'), 'rb'))
def cal(p, side):
    t = CAL[side]['table']; return float(np.interp(p, t.p_mean, t.hit_mono))
BASE_P = {'buy': CAL['yb']['base'], 'sell': CAL['ys']['base']}
# 2) today's rule signals
SG = pd.read_csv(os.path.join(ST, 'signals.csv'))
SG = SG[SG.time_utc.astype(str).str[:10] == day]
SG = SG[~SG.rule.astype(str).str.startswith('5%')]
SG['side'] = np.where(SG.strength.astype(str).str.startswith('فروش'), 'sell', 'buy')
DL = os.path.join(ST, 'decisions_10m.csv')
if os.path.exists(DL):
    D = pd.read_csv(DL); D = D[(D.day == day) & D.decision.astype(str).str.startswith('تأیید')]
    D = D.assign(side=np.where(D.strength.astype(str).str.startswith('فروش'), 'sell', 'buy'), note=D.decision)
    SG = pd.concat([SG[['sym', 'rule', 'side', 'note']], D[['sym', 'rule', 'side', 'note']]])
SG = SG.drop_duplicates(['sym', 'rule', 'side'])
SM = pd.read_csv(os.path.join(ST, 'symbol_methods.csv')).set_index(['sym', 'rule'])
RULE_ALL = pd.read_csv(os.path.join(ST, 'symbol_methods.csv')).groupby('rule').apply(lambda d: (d.hist_hit * d.n).sum() / max(d.n.sum(), 1))
LQ = pd.read_csv(os.path.join(ST, 'tse-data', 'data', 'latest.csv'), low_memory=False)
LQ['sym'] = LQ['نماد'].map(live.norm); QUEUE = LQ.drop_duplicates('sym').set_index('sym')['صف'].fillna('') if 'صف' in LQ else pd.Series(dtype=str)
DIV = pd.read_csv(os.path.join(ST, 'divergence.csv')); DIV = DIV[DIV.date.astype(str) == day].set_index('sym')['divergence']
# 3) per-symbol verdict
rows = []
for k, x in X.items():
    if x.index[-1] != DAY: continue
    sym = live.norm(k); r = x.iloc[-1]
    ev = {'buy': [], 'sell': []}; why = {'buy': [], 'sell': []}
    if sym in S.index:
        pb, ps = cal(S.at[sym, 'p_yb'], 'yb'), cal(S.at[sym, 'p_ys'], 'ys')
        ev['buy'].append((pb, 2.0)); ev['sell'].append((ps, 2.0))
    else: pb = ps = np.nan
    for g in SG[SG.sym == sym].itertuples():
        if (sym, g.rule) in SM.index:
            m = SM.loc[(sym, g.rule)]; m = m.iloc[0] if isinstance(m, pd.DataFrame) else m
            nh, nl = float(m.n), float(m.n_live) if pd.notna(m.n_live) else 0.0
            kh = float(m.hist_hit) * nh if pd.notna(m.hist_hit) else 0.0; kl = float(m.live_hit) * nl if pd.notna(m.live_hit) else 0.0
            # same blend as symbol_methods (live x6 of history) but shrunk toward the market base rate of this side, not 50%
            sc = (0.5 * kh + 3 * kl + 2 * BASE_P[g.side]) / (0.5 * nh + 3 * nl + 2); n = nh + nl
        else:
            continue   # no record of this rule on this symbol: no evidence
        if np.isnan(sc): continue
        w = 1.0 if n >= 10 else 0.5
        ev[g.side].append((sc, w)); why[g.side].append(f'{g.rule} ({sc:.0%}' + (f'، {int(n)} نمونه)' if n else '، میانگین بازار)'))
    if sym in DIV.index:
        why['sell'].append(str(DIV[sym]))
    conf = {s: (sum(p * w for p, w in ev[s]) / sum(w for _, w in ev[s])) if ev[s] else np.nan for s in ev}
    edge = {s: conf[s] - BASE_P[s] if not np.isnan(conf[s]) else -1 for s in conf}
    # direction must be the more likely of the two outcomes AND beat that side's market base rate by 5 points
    cb, cs = (conf['buy'] if not np.isnan(conf['buy']) else -1), (conf['sell'] if not np.isnan(conf['sell']) else -1)
    side = 'buy' if cb >= cs else 'sell'
    verdict = {'buy': 'خرید', 'sell': 'فروش'}[side] if edge[side] >= 0.05 else 'خنثی'
    liquid = (r.value >= 2e10)
    tier = '' if verdict == 'خنثی' else ('قوی' if edge[side] >= 0.10 else 'متوسط')
    rows.append({'نماد': sym, 'سیگنال': verdict, 'قدرت سیگنال': tier, 'اطمینان': conf[side] if verdict != 'خنثی' else np.nan,
                 'احتمال خرید (۵٪ قبل از −۵٪، ۱۰ روز)': conf['buy'], 'احتمال فروش (−۵٪ قبل از +۵٪، ۱۰ روز)': conf['sell'],
                 'برتری نسبت به میانگین بازار': edge[side] if verdict != 'خنثی' else np.nan,
                 'مدل ۵٪ خرید (کالیبره)': pb, 'مدل ۵٪ فروش (کالیبره)': ps,
                 'قیمت پایانی': r.c, 'تغییر امروز': r.rawpct, 'ارزش معاملات (میلیارد تومان)': r.value / 1e10,
                 'قدرت خریدار حقیقی': r.bp, 'ورود پول حقیقی / ارزش': r.nrp,
                 'نقدشوندگی': 'کافی' if liquid else 'کم (زیر ۲ میلیارد تومان)',
                 'وضعیت صف در پایان روز': QUEUE.get(sym, '') or '—',
                 'دلایل خرید': ' | '.join(why['buy']), 'دلایل فروش': ' | '.join(why['sell'])})
R = pd.DataFrame(rows)
R['_o'] = R['سیگنال'].map({'خرید': 0, 'فروش': 1, 'خنثی': 2})
R = R.sort_values(['_o', 'برتری نسبت به میانگین بازار'], ascending=[True, False]).drop(columns='_o')
# 4) other sheets
dec = pd.read_csv(DL) if os.path.exists(DL) else pd.DataFrame()
dec_today = dec[dec.day == day][['time', 'sym', 'rule', 'strength', 'last', 'lpct', 'bp', 'nrp', 'pace', 'inrp', 'decision']] if len(dec) else dec
dec_today = dec_today.rename(columns={'time': 'ساعت', 'sym': 'نماد', 'rule': 'قاعده', 'strength': 'جهت', 'last': 'قیمت', 'lpct': 'تغییر٪', 'bp': 'قدرت خریدار', 'nrp': 'ورود حقیقی/ارزش', 'pace': 'سرعت ارزش', 'inrp': 'ورود حقیقی ۳۰ دقیقه/ارزش', 'decision': 'تصمیم'})
EV = pd.read_csv(os.path.join(ST, 'decision_eval.csv')) if os.path.exists(os.path.join(ST, 'decision_eval.csv')) else pd.DataFrame()
if len(EV):
    EV['dec'] = EV.decision.str.split(':').str[0]
    acc = EV.groupby(['day', 'rule', 'dec']).agg(n=('sym', 'size'), r1=('ok_r1', 'mean'), g1=('g_r1', 'mean'), r3=('ok_r3', 'mean'), g3=('g_r3', 'mean'), r5=('ok_r5', 'mean'),
                                                 hit5=('tp5_10d', lambda s: (s == 'درست').sum()), miss5=('tp5_10d', lambda s: (s == 'غلط').sum()), open5=('tp5_10d', lambda s: (s == 'باز').sum())).reset_index()
    acc.columns = ['روز سیگنال', 'قاعده', 'تصمیم', 'تعداد', 'درست روز بعد', 'سود میانگین روز بعد', 'درست ۳ روز', 'سود میانگین ۳ روز', 'درست ۵ روز', '۵٪ درست', '۵٪ غلط', 'هنوز باز']
else: acc = pd.DataFrame()
cal_rows = []
for side, nm in [('yb', 'خرید'), ('ys', 'فروش')]:
    t = CAL[side]['table']
    for q in t.itertuples(): cal_rows.append({'سمت': nm, 'احتمال خام از': q.p_lo, 'تا': q.p_hi, 'نرخ واقعی (۱۴۰۴–۱۴۰۵)': q.hit_mono, 'تعداد نمونه': q.n})
CALT = pd.DataFrame(cal_rows)
guide = [
    ('این فایل چیست', f'نتیجهٔ بررسی پایان روز {day} برای همهٔ نمادهایی که امروز معامله شدند ({len(R)} نماد).'),
    ('سیگنال', 'خرید یا فروش وقتی آن سمت محتمل‌تر از سمت دیگر باشد و دست‌کم ۵ واحد درصد از میانگین بازار همان سمت بهتر باشد؛ وگرنه خنثی.'),
    ('قدرت سیگنال', 'قوی: اطمینان دست‌کم ۱۰ واحد درصد بالاتر از میانگین بازار. متوسط: ۵ تا ۱۰ واحد.'),
    ('اطمینان', 'احتمال برآوردی این که قیمت در ۱۰ روز معاملاتی آینده اول ۵٪ در جهت سیگنال حرکت کند، قبل از این که ۵٪ خلاف آن برود. مبنای قیمت: قیمت پایانی امروز.'),
    ('میانگین بازار', f'در ۱۴۰۴–۱۴۰۵ این اتفاق برای خرید در {BASE_P["buy"]:.0%} و برای فروش در {BASE_P["sell"]:.0%} نمادها رخ داده است. اطمینان را با این عددها مقایسه کنید، نه با ۵۰٪.'),
    ('برتری نسبت به میانگین بازار', 'اطمینان منهای میانگین بازار همان سمت. ستون مرتب‌سازی اصلی است.'),
    ('مدل ۵٪ (کالیبره)', 'احتمال مدل یادگیری ماشین، تبدیل‌شده به نرخ واقعی که روی دادهٔ ۱۴۰۴–۱۴۰۵ (خارج از دادهٔ آموزش) دیده شد. جدول تبدیل در برگهٔ «کالیبراسیون» است.'),
    ('دلایل', 'قاعده‌هایی که امروز برای این نماد فعال شدند، با نرخ موفقیت همان قاعده روی همین نماد (ترکیب تاریخچه و نتایج زنده، با وزن بیشتر روی زنده). نرخ هر قاعده به سمت میانگین بازار کشیده شده تا نمونهٔ کم باعث اطمینان کاذب نشود.'),
    ('ترکیب', 'اطمینان = میانگین وزنی: مدل ۵٪ با وزن ۲، هر قاعده با وزن ۱ (یا ۰٫۵ اگر کمتر از ۱۰ نمونه دارد).'),
    ('وضعیت صف در پایان روز', 'اگر نماد در صف خرید بسته شده، خرید با قیمت پایانی عملاً ممکن نبوده و فردا معمولاً با قیمت بالاتر باز می‌شود؛ برای فروش در صف فروش هم همین‌طور.'),
    ('سیگنال فروش', 'در روزهای رالی، سیگنال‌های فروش معمولاً به‌طور مطلق بالا می‌روند ولی کمتر از بازار؛ ستون برتری را به معنی «ضعیف‌تر از بازار» هم بخوانید.'),
    ('محدودیت‌ها', 'بیشترین احتمال واقعی مدل خارج از نمونه حدود ۵۵٪ برای خرید و ۵۳٪ برای فروش بوده است؛ اطمینان بالاتر از این‌ها بیشتر از قاعده‌هاست و با نمونهٔ کم کمتر قابل اتکاست. نمادهای کم‌نقد یا قفل در صف ممکن است عملاً قابل معامله نباشند. این برآورد آماری است، نه توصیهٔ سرمایه‌گذاری.'),
    ('برگه‌ها', 'سیگنال‌ها | تصمیم‌های ۱۰ دقیقه‌ای امروز | دقت روزهای قبل | کالیبراسیون | راهنما'),
]
out = os.path.join(ST, 'reports'); os.makedirs(out, exist_ok=True)
path = os.path.join(out, f'signals_{day}.xlsx')
wb = Workbook(); wb.remove(wb.active)
HDR = PatternFill('solid', fgColor='1F3864'); FB = PatternFill('solid', fgColor='E2EFDA'); FS = PatternFill('solid', fgColor='FCE4D6')
PCT = {'اطمینان', 'احتمال خرید (۵٪ قبل از −۵٪، ۱۰ روز)', 'احتمال فروش (−۵٪ قبل از +۵٪، ۱۰ روز)', 'برتری نسبت به میانگین بازار', 'مدل ۵٪ خرید (کالیبره)', 'مدل ۵٪ فروش (کالیبره)',
       'تغییر امروز', 'ورود پول حقیقی / ارزش', 'درست روز بعد', 'سود میانگین روز بعد', 'درست ۳ روز', 'سود میانگین ۳ روز', 'درست ۵ روز', 'نرخ واقعی (۱۴۰۴–۱۴۰۵)', 'احتمال خام از', 'تا', 'ورود حقیقی/ارزش', 'ورود حقیقی ۳۰ دقیقه/ارزش'}
def sheet(name, df, widths=None, color_col=None):
    ws = wb.create_sheet(name); ws.sheet_view.rightToLeft = True
    ws.append(list(df.columns))
    for c in ws[1]: c.font = Font(name='Arial', bold=True, color='FFFFFF'); c.fill = HDR; c.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
    for row in df.itertuples(index=False):
        ws.append([None if (isinstance(v, float) and np.isnan(v)) else v for v in row])
    for j, col in enumerate(df.columns, 1):
        L = get_column_letter(j)
        for c in ws[L][1:]:
            c.font = Font(name='Arial')
            if col in PCT: c.number_format = '0.0%'
            elif col in ('قیمت پایانی', 'قیمت'): c.number_format = '#,##0'
            elif df[col].dtype.kind == 'f': c.number_format = '#,##0.00'
        ws.column_dimensions[L].width = (widths or {}).get(col, 14 if col not in ('دلایل خرید', 'دلایل فروش') else 70)
    if color_col and color_col in df.columns:
        j = list(df.columns).index(color_col) + 1
        for i in range(2, ws.max_row + 1):
            v = ws.cell(i, j).value
            if v in ('خرید', 'فروش'):
                for c in ws[i]: c.fill = FB if v == 'خرید' else FS
    ws.freeze_panes = 'B2'; ws.row_dimensions[1].height = 45
    if ws.max_row > 1: ws.auto_filter.ref = ws.dimensions
    return ws
sheet('سیگنال‌ها', R, color_col='سیگنال')
sheet('تصمیم‌های ۱۰ دقیقه‌ای', dec_today)
sheet('دقت روزهای قبل', acc)
sheet('کالیبراسیون', CALT)
ws = sheet('راهنما', pd.DataFrame(guide, columns=['موضوع', 'توضیح']), widths={'موضوع': 26, 'توضیح': 120})
for c in ws['B'][1:]: c.alignment = Alignment(wrap_text=True, vertical='top')
wb.save(path)
print(path); print(R['سیگنال'].value_counts().to_string())
print(R.head(12)[['نماد', 'سیگنال', 'اطمینان', 'برتری نسبت به میانگین بازار', 'نقدشوندگی']].round(3).to_string(index=False))
print(R[R['سیگنال'] == 'فروش'].head(12)[['نماد', 'سیگنال', 'اطمینان', 'برتری نسبت به میانگین بازار', 'نقدشوندگی']].round(3).to_string(index=False))
